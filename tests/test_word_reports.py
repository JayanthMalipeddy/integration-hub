"""Word exports must contain the selected evidence and retain authentication/redaction."""
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from backend.app.main import create_app
from backend.app.models import Finding
from backend.app.reports import DOCX_MEDIA_TYPE, investigation_docx
from tests.test_workflow import FixtureLLM, build, state


def text_from_docx(content):
    doc = Document(BytesIO(content))
    return "\n".join([p.text for p in doc.paragraphs] +
                     [cell.text for table in doc.tables for row in table.rows for cell in row.cells])


def test_word_download_is_valid_attachment_for_selected_run(tmp_path):
    settings, store, providers, _ = build(tmp_path)
    selected = state()
    selected.update(status="awaiting_approval", finding=FixtureLLM().generate([], Finding.model_json_schema())["content"],
                    calls=[{"provider":"fixture", "input_tokens":42}], validation={"passed": True})
    store.put("run", "local", selected["id"], selected)
    with TestClient(create_app(settings, providers)) as client:
        response = client.get(f"/api/runs/{selected['id']}/report.docx")
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_MEDIA_TYPE
    assert response.headers["content-disposition"] == f'attachment; filename="investigation-{selected["id"][:8]}.docx"'
    with ZipFile(BytesIO(response.content)) as archive:
        assert "word/document.xml" in archive.namelist()
    text = text_from_docx(response.content)
    for expected in [selected["incident"]["title"], selected["id"], selected["finding"]["hypothesis"],
                     selected["finding"]["uncertainty"], selected["finding"]["citations"][1]["quote"],
                     selected["finding"]["checks"][0], selected["finding"]["recovery"][0], '"input_tokens": 42']:
        assert expected in text
    assert "Unverified draft" not in text


def test_word_export_requires_auth_and_sanitizes_legacy_records(tmp_path):
    settings, store, providers, _ = build(tmp_path)
    settings.require_auth = True
    settings.api_token = "test-word-export-credential"
    selected = state()
    selected.update(status="failed", error="Workflow stopped", finding={"hypothesis":"password=secret-for-export"})
    # Bypass current write redaction to simulate a legacy saved record.
    store.db.execute("INSERT INTO records VALUES (?,?,?,?)", ("run", "local", selected["id"], json.dumps(selected)))
    store.db.commit()
    with TestClient(create_app(settings, providers)) as client:
        url = f"/api/runs/{selected['id']}/report.docx"
        assert client.get(url).status_code == 401
        client.headers["Authorization"] = "Bearer " + settings.api_token
        response = client.get(url)
        assert response.status_code == 200
        text = text_from_docx(response.content)
        assert "secret-for-export" not in text and "[REDACTED]" in text
        assert "Unverified draft" in text and "Workflow stopped" in text
        assert client.get("/api/runs/missing/report.docx").status_code == 404


def test_empty_finding_unicode_and_log_control_characters_export():
    selected = state()
    selected["incident"]["title"] = "Incident <XML> & Oracle ? ???????\x00\x1b"
    selected.update(status="needs_human", finding={})
    text = text_from_docx(investigation_docx(selected))
    assert "<XML> & Oracle ? ???????" in text
    assert "No supported finding" in text and "No supporting quotations recorded." in text
