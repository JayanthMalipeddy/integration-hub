"""Portable investigation reports generated in memory from redacted run data."""
from io import BytesIO
import json
import re

from docx import Document
from docx.shared import Inches, Pt

from .security import redact_data

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def xml_text(value):
    # Logs can contain terminal control characters that XML/Word cannot represent.
    text = str(value)
    return "".join(c for c in text if c in "\t\n\r" or
                   0x20 <= ord(c) <= 0xD7FF or 0xE000 <= ord(c) <= 0xFFFD or
                   0x10000 <= ord(c) <= 0x10FFFF)


def report_filename(identifier):
    safe = re.sub(r"[^A-Za-z0-9_-]", "", str(identifier))[:8] or "report"
    return f"investigation-{safe}.docx"


def investigation_docx(run):
    data = redact_data(run)
    finding = data.get("finding") or {}
    document = Document()
    document.core_properties.title = xml_text(data["incident"]["title"])
    document.core_properties.author = "Integration Hub"
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(xml_text(data["incident"]["title"]), 0)
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, value in [("Investigation", data.get("id", "")), ("Status", data["status"]),
                         ("Affected service", data["incident"].get("service", "Not specified")),
                         ("Created", data.get("created_at", "Not recorded")),
                         ("Profile", data.get("profile", "Not recorded")),
                         ("Confidence", finding.get("confidence", "Not assessed"))]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, xml_text(value)
    if data.get("error"):
        document.add_paragraph("Investigation notice: " + xml_text(data["error"]))
    document.add_heading("Hypothesis (requires operator verification)", 1)
    document.add_paragraph(xml_text(finding.get("hypothesis", "No supported finding")))
    if not (data.get("validation") or {}).get("passed") and finding:
        document.add_paragraph("Unverified draft: evidence validation has not passed.")
    if data.get("observations"):
        document.add_heading("Observed evidence (operator supplied)", 1)
        for fact in data["observations"]:
            document.add_paragraph(xml_text(fact["line_id"] + ": " + fact["quote"]))
    if data.get("questions"):
        document.add_heading("Requested follow-up evidence", 1)
        for question in data["questions"]:
            document.add_paragraph(xml_text(question), style="List Bullet")
    document.add_paragraph("Evidence revision: " + str(data.get("evidence_revision", 1)) + ". Root cause has not been independently verified.")
    document.add_heading("What remains uncertain", 1)
    document.add_paragraph(xml_text(finding.get("uncertainty", "Not assessed")))
    document.add_heading("Evidence", 1)
    citations = finding.get("citations") or []
    for citation in citations:
        document.add_paragraph(xml_text(citation["source_id"]), style="Heading 2")
        document.add_paragraph(xml_text(citation["quote"]))
    if not citations:
        document.add_paragraph("No supporting quotations recorded.")
    for heading, key in [("Recommended checks", "checks"), ("Proposed recovery", "recovery")]:
        document.add_heading(heading, 1)
        for item in finding.get(key) or []:
            document.add_paragraph(xml_text(item), style="List Bullet")
        if not finding.get(key):
            document.add_paragraph("None recorded.")
    document.add_heading("Provider provenance", 1)
    document.add_paragraph("Recorded model calls; these measurements are not a diagnosis accuracy score.")
    document.add_paragraph(xml_text(json.dumps(data.get("calls", []), indent=2, ensure_ascii=False)))
    output = BytesIO()
    document.save(output)
    return output.getvalue()
