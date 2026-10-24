"""Build the local Word documentation pack from its Markdown sources.

Documentation-only dependency: python-docx. These are already
installed with requirements-docs.txt; not required by the application.
"""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
DIAGRAMS = DOCS / 'diagrams'


def architecture():
    """Standalone editable SVG; Word receives a native table-based diagram."""
    from xml.sax.saxutils import escape
    DIAGRAMS.mkdir(exist_ok=True)
    nodes = [
        (30,80,'Browser workspace','React / TypeScript :5180'),
        (380,80,'API boundary','FastAPI :8010'),
        (730,80,'Canonical records','SQLite docs, runs, traces, evals'),
        (30,240,'Provider factory','AI_PROFILE selects adapters'),
        (380,240,'LangGraph workflow','Plan, retrieve, analyze; bounded correction'),
        (730,240,'Durable checkpoints','SQLite thread state and interrupts'),
        (30,400,'Verified local providers','Ollama Qwen + embeddings; Chroma'),
        (380,400,'Human approval gate','Review evidence and uncertainty'),
        (730,400,'Human handoff / retry','Insufficient evidence or saved failure'),
        (30,560,'Optional cloud adapters','Configured and gated; not live-verified'),
        (380,560,'Approved local outcome','Ticket and reviewed hypothesis'),
        (730,560,'Measured evaluations','Actual cases feed dashboard records'),
    ]
    parts=['<svg xmlns="http://www.w3.org/2000/svg" width="1060" height="730" viewBox="0 0 1060 730">',
           '<rect width="1060" height="730" fill="#f7fafc"/>',
           '<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0 0 L8 4 L0 8 Z" fill="#087d70"/></marker></defs>',
           '<text x="30" y="36" font-family="Arial" font-size="25" font-weight="bold" fill="#183049">Integration Hub - System architecture</text>']
    for x,y,title,sub in nodes:
        parts.append(f'<rect x="{x}" y="{y}" width="290" height="90" rx="10" fill="#e6f5f1" stroke="#8aa9b6"/>')
        parts.append(f'<text x="{x+145}" y="{y+34}" text-anchor="middle" font-family="Arial" font-size="17" font-weight="bold" fill="#183049">{escape(title)}</text>')
        parts.append(f'<text x="{x+145}" y="{y+62}" text-anchor="middle" font-family="Arial" font-size="12" fill="#52677b">{escape(sub)}</text>')
    for x1,y1,x2,y2 in [(320,125,375,125),(670,125,725,125),(525,170,525,235),(380,285,325,285),(670,285,725,285),(175,330,175,395),(525,330,525,395),(670,325,725,410),(525,490,525,555)]:
        parts.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#087d70" stroke-width="2" marker-end="url(#arrow)"/>')
    parts.append('<path d="M30 285 L12 285 L12 605 L25 605" fill="none" stroke="#087d70" stroke-dasharray="6 4" marker-end="url(#arrow)"/>')
    parts.append('<text x="30" y="700" font-family="Arial" font-size="13" fill="#52677b">Solid: local path. Dashed: optional cloud configuration. Approval gates tickets and reviewed memory only.</text></svg>')
    (DIAGRAMS/'system_architecture.svg').write_text(''.join(parts),encoding='utf-8')


def word_architecture(doc):
    """Editable diagram rendered entirely with Word-native cells and arrows."""
    rows = [
        ['Browser workspace\nReact / TypeScript :5180', '\u2192', 'API boundary\nFastAPI :8010', '\u2194', 'Canonical records\nSQLite documents, runs, traces, evals'],
        ['', '', '\u2193', '', ''],
        ['Provider factory\nAI_PROFILE selects adapters', '\u2190', 'LangGraph workflow\nPlan > tools > analyze > verify\nAt most two corrections', '\u2194', 'Durable checkpoints\nSQLite state and interrupts'],
        ['\u2193 local', '', '\u2193 valid evidence', '', 'Invalid evidence / failure \u2193'],
        ['Verified local providers\nOllama: Qwen + embeddings\nChroma: semantic retrieval', '', 'Human approval gate\nReview hypothesis, evidence\nand uncertainty', '', 'Human handoff / retry\nEvidence gaps need review\nSaved failures can resume'],
        ['', '', '\u2193 approved', '', ''],
        ['Optional cloud adapters\nNebius / Pinecone / Mem0\nExplicit trace, search and voice\nConfiguration + consent required', '', 'Approved local outcome\nRecovery ticket + reviewed memory\nNo production remediation', '', 'Measured evaluations\nExecuted cases > saved results\nSame records feed dashboard'],
    ]
    diagram=doc.add_table(rows=0,cols=5)
    diagram.autofit=False
    widths=[3.05,.45,3.05,.45,3.05]
    for row_index,values in enumerate(rows):
        cells=diagram.add_row().cells
        for column,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Inches(widths[column])
            para=cell.paragraphs[0];para.alignment=WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before=Pt(9 if row_index%2==0 else 3)
            para.paragraph_format.space_after=Pt(9 if row_index%2==0 else 3)
            for n,line in enumerate(value.split('\n')):
                if n:para.add_run().add_break()
                r=para.add_run(line);r.font.size=Pt(10 if row_index%2==0 else 9)
                r.bold=(n==0);r.font.color.rgb=RGBColor.from_string('183049')
            if row_index%2==0 and column%2==0:
                pr=cell._tc.get_or_add_tcPr();shading=OxmlElement('w:shd')
                shading.set(qn('w:fill'),'E6F5F1' if (row_index,column)!=(4,2) else 'FFF4DC');pr.append(shading)
        trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    doc.add_paragraph('Optional cloud adapters connect through the provider factory. They are not part of the verified local execution path.',style='Caption')


def field(paragraph,code):
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),code);paragraph._p.append(f)


def layout(section,landscape=False):
    section.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width=Inches(11.69 if landscape else 8.27)
    section.page_height=Inches(8.27 if landscape else 11.69)
    section.top_margin=section.bottom_margin=Inches(.65)
    section.left_margin=section.right_margin=Inches(.72)
    section.header_distance=section.footer_distance=Inches(.28)


def text_runs(paragraph,text):
    # Keep filenames and links readable in Word without raw Markdown syntax.
    text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',text)
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)',text):
        if part.startswith('**'):
            r=paragraph.add_run(part[2:-2]);r.bold=True
        elif part.startswith('`'):
            r=paragraph.add_run(part[1:-1]);r.font.name='Consolas';r.font.size=Pt(9)
        else:paragraph.add_run(part)


def table(doc,lines):
    rows=[[s.strip() for s in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[:\- ]+',s) for s in r)]
    t=doc.add_table(rows=1,cols=len(rows[0]));t.style='Light Shading Accent 1'
    for index,row in enumerate(rows):
        cells=t.rows[0].cells if index==0 else t.add_row().cells
        for c,content in zip(cells,row):
            text_runs(c.paragraphs[0],content)
            for run in c.paragraphs[0].runs:run.font.size=Pt(9)
            c.paragraphs[0].paragraph_format.space_after=Pt(4)
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        element=OxmlElement('w:cantSplit');trpr.append(element)
        if index==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for c in cells:
                for r in c.paragraphs[0].runs:r.bold=True
    doc.add_paragraph()


def build(source,filename,title):
    doc=Document();layout(doc.sections[0])
    styles=doc.styles
    styles['Normal'].font.name='Calibri';styles['Normal'].font.size=Pt(10.5)
    styles['Normal'].paragraph_format.space_after=Pt(7)
    styles['Normal'].paragraph_format.line_spacing=1.1
    for name in ['Title','Heading 1','Heading 2','Heading 3']:
        styles[name].font.name='Calibri';styles[name].font.color.rgb=RGBColor.from_string('183049')
    styles['Title'].font.size=Pt(30)
    styles['Heading 1'].font.size=Pt(17)
    styles['Heading 2'].font.size=Pt(13)
    header=doc.sections[0].header.paragraphs[0]
    header.text='INTEGRATION HUB  |  WEEK 3'
    header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor.from_string('52677B')
    footer=doc.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(title+'  |  Page ');field(footer,'PAGE')
    for r in footer.runs:r.font.size=Pt(8)
    doc.core_properties.title='Integration Hub - '+title
    doc.core_properties.subject='Week 3 implemented application documentation'
    doc.core_properties.author='Jayanth'
    lines=(DOCS/source).read_text(encoding='utf-8').splitlines()
    i=0; code=False;code_type='';landscape=False
    while i<len(lines):
        line=lines[i]
        if line.startswith('```'):
            if not code:code_type=line[3:].strip()
            code=not code;i+=1;continue
        if code:
            if code_type!='mermaid':
                p=doc.add_paragraph();r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(8.5)
                p.paragraph_format.space_after=Pt(1)
            i+=1;continue
        if line.startswith('## 2. System architecture'):
            layout(doc.add_section(WD_SECTION_START.NEW_PAGE),True);landscape=True
        elif landscape and line.startswith('## 3.'):
            layout(doc.add_section(WD_SECTION_START.NEW_PAGE));landscape=False
        if line.startswith('!['):
            word_architecture(doc)
            i+=1
            # The next paragraph is the figure caption; keep it compact on the landscape page.
            while i<len(lines) and not lines[i]:i+=1
            if i<len(lines) and not lines[i].startswith('#'):
                p=doc.add_paragraph(lines[i]);p.style='Caption'
                for r in p.runs:r.font.size=Pt(9)
                i+=1
            continue
        if line.startswith('|'):
            group=[]
            while i<len(lines) and lines[i].startswith('|'):group.append(lines[i]);i+=1
            table(doc,group);continue
        if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):doc.add_heading(line[3:],level=1)
        elif line.startswith('### '):doc.add_heading(line[4:],level=2)
        elif line.startswith('- '):text_runs(doc.add_paragraph(style='List Bullet'),line[2:])
        elif re.match(r'^\d+\. ',line):text_runs(doc.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',line))
        elif line:
            p=doc.add_paragraph();text_runs(p,line)
            if re.match(r'^(SHOW|DO|SAY|CHECK):',line):
                for r in p.runs:
                    if line.startswith('SAY:'):r.font.size=Pt(11)
                if not line.startswith('SAY:'):p.paragraph_format.space_after=Pt(6)
        i+=1
    output=DOCS/filename;doc.save(output)
    # Reopen each file and inspect package structure instead of assuming generation worked.
    checked=Document(output)
    assert len(checked.paragraphs)>15 and len(checked.sections)>=1
    if 'Technical' in filename:assert len(checked.tables)>=5 and len(checked.sections)==3
    print(f'{output.name}: {output.stat().st_size:,} bytes, {len(checked.paragraphs)} paragraphs, {len(checked.tables)} tables')


if __name__=='__main__':
    architecture()
    build('TECHNICAL_SPECIFICATION.md','Integration_Hub_Technical_Specification_Quality.docx','Technical specification')
    build('PROJECT_SPECIFICATION.md','Integration_Hub_Project_Specification_Quality.docx','Project specification')
    build('VIDEO_DEMO_SCRIPT.md','Integration_Hub_Video_Demo_Script.docx','Video demonstration script')

    if (DOCS / "QUALITY_VERIFICATION.md").exists():
        build("QUALITY_VERIFICATION.md", "Integration_Hub_Quality_Verification.docx", "Quality verification")
